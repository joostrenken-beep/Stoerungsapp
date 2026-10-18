"""Gemeinsame Import-Logik: XML (ISYBAU / DWA-M 150), CSV und Excel.

Alle Importe sind wiederholbar: Schächte, Haltungen und Anlagen werden über
ihre Bezeichnung abgeglichen und aktualisiert. Eine Inspektion desselben
Objekts am selben Datum wird beim erneuten Import ersetzt.
"""
import csv
import io
import os
import shutil
import traceback
from datetime import datetime

from flask import current_app

from ..dwa149 import gesamtklasse, klassen_vorschlag, zustand_aktualisieren
from ..extensions import db
from ..geo import nach_wgs84
from ..models import Anlage, Haltung, ImportProtokoll, Inspektion, Knoten, Schaden
from . import isybau

ENDUNGEN = (".xml", ".csv", ".xlsx")


class ImportFehler(Exception):
    pass


def _setzen(obj, werte, felder):
    for feld in felder:
        if feld in werte and werte[feld] is not None:
            setattr(obj, feld, werte[feld])


KNOTEN_FELDER = ["knotenart", "kanalart", "ortsteil", "strasse", "rechtswert", "hochwert", "lat", "lon",
                 "deckelhoehe", "sohlhoehe", "material", "baujahr", "status", "bemerkung"]
HALTUNG_FELDER = ["kanalart", "funktion", "ortsteil", "strasse", "laenge", "dn", "profilbreite", "profilart",
                  "material", "baujahr", "sohle_zulauf", "sohle_ablauf", "status", "geometrie", "bemerkung"]
ANLAGE_FELDER = ["name", "anlagentyp", "kanalart", "ortsteil", "adresse", "lat", "lon", "baujahr",
                 "volumen_m3", "vorfluter", "status", "bemerkung"]
SCHADEN_FELDER = ["kode", "char1", "char2", "quant1", "quant2", "station", "station_bis", "uhr_von", "uhr_bis",
                  "verbindung", "klasse_d", "klasse_s", "klasse_b", "schadensklasse", "foto", "bemerkung"]


def speichern(daten, quelle, protokoll):
    """Schreibt die vom Parser gelieferten Daten in die Datenbank."""
    meldungen = []
    knoten_cache = {k.bezeichnung: k for k in Knoten.query.all()}
    for werte in daten.get("knoten", []):
        k = knoten_cache.get(werte["bezeichnung"])
        if k is None:
            k = Knoten(bezeichnung=werte["bezeichnung"])
            db.session.add(k)
            knoten_cache[k.bezeichnung] = k
        _setzen(k, werte, KNOTEN_FELDER)
        k.quelle = quelle
        protokoll.knoten += 1
    db.session.flush()

    haltung_cache = {h.bezeichnung: h for h in Haltung.query.all()}
    for werte in daten.get("haltungen", []):
        h = haltung_cache.get(werte["bezeichnung"])
        if h is None:
            h = Haltung(bezeichnung=werte["bezeichnung"])
            db.session.add(h)
            haltung_cache[h.bezeichnung] = h
        _setzen(h, werte, HALTUNG_FELDER)
        for seite in ("von", "bis"):
            name = werte.get(seite)
            if name:
                if name in knoten_cache:
                    setattr(h, f"{seite}_knoten", knoten_cache[name])
                else:
                    meldungen.append(f"Haltung {h.bezeichnung}: Knoten {name} unbekannt")
        h.quelle = quelle
        protokoll.haltungen += 1
    db.session.flush()

    for werte in daten.get("anlagen", []):
        a = Anlage.query.filter_by(bezeichnung=werte["bezeichnung"]).first()
        if a is None:
            a = Anlage(bezeichnung=werte["bezeichnung"])
            db.session.add(a)
        _setzen(a, werte, ANLAGE_FELDER)
        protokoll.anlagen += 1

    betroffen = set()
    for insp in daten.get("inspektionen", []):
        cache = haltung_cache if insp["ziel_typ"] == "haltung" else knoten_cache
        ziel = cache.get(insp["objekt"])
        if ziel is None:
            meldungen.append(f"Inspektion: Objekt {insp['objekt']} nicht in der Datenbank – übersprungen")
            continue
        vorhanden = Inspektion.query.filter_by(ziel_typ=insp["ziel_typ"], ziel_id=ziel.id, datum=insp["datum"]).all()
        for alt in vorhanden:
            db.session.delete(alt)
        i = Inspektion(ziel_typ=insp["ziel_typ"], ziel_id=ziel.id, datum=insp["datum"],
                       untersucher=insp.get("untersucher"), firma=insp.get("firma"),
                       richtung=insp.get("richtung"), inspektionslaenge=insp.get("inspektionslaenge"),
                       zustandsklasse=insp.get("zustandsklasse"), quelle=quelle)
        db.session.add(i)
        for s in insp["schaeden"]:
            schaden = Schaden(ziel_typ=insp["ziel_typ"], ziel_id=ziel.id, datum=insp["datum"],
                              erstellt_von=f"Import {quelle}")
            _setzen(schaden, s, SCHADEN_FELDER)
            i.schaeden.append(schaden)
            protokoll.schaeden += 1
        protokoll.inspektionen += 1
        betroffen.add(ziel)
    db.session.flush()
    for ziel in betroffen:
        zustand_aktualisieren(ziel)
    return meldungen


# ---------------------------------------------------------------- Tabellen

def _tabelle_lesen(pfad):
    if pfad.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(pfad, read_only=True, data_only=True)
        zeilen = list(wb.active.iter_rows(values_only=True))
        if not zeilen:
            return []
        kopf = [str(c or "").strip().lower() for c in zeilen[0]]
        return [dict(zip(kopf, z)) for z in zeilen[1:] if any(v is not None for v in z)]
    with open(pfad, "rb") as f:
        roh = f.read()
    try:
        text = roh.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = roh.decode("cp1252")
    trenner = ";" if text.split("\n", 1)[0].count(";") >= text.split("\n", 1)[0].count(",") else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=trenner)
    return [{(k or "").strip().lower(): v for k, v in z.items()} for z in reader]


def _wert(v, typ):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if typ == "float":
        return isybau.zahl(v)
    if typ == "int":
        return isybau.ganzzahl(v)
    if typ == "date":
        return v.date() if isinstance(v, datetime) else isybau.datum(str(v))
    if typ == "bool":
        return str(v).strip().lower() in ("ja", "j", "1", "true", "x")
    return str(v).strip()


TABELLEN_TYPEN = {
    "rechtswert": "float", "hochwert": "float", "lat": "float", "lon": "float",
    "deckelhoehe": "float", "sohlhoehe": "float", "baujahr": "int", "laenge": "float", "dn": "int",
    "profilbreite": "int", "sohle_zulauf": "float", "sohle_ablauf": "float", "volumen_m3": "float",
    "station": "float", "station_bis": "float", "uhr_von": "int", "uhr_bis": "int", "verbindung": "bool",
    "klasse_d": "int", "klasse_s": "int", "klasse_b": "int", "schadensklasse": "int", "datum": "date",
}


def tabelle_lesen(pfad, zone=32):
    """Erkennt anhand der Spaltenköpfe, ob Schächte, Haltungen, Schäden oder Anlagen enthalten sind."""
    zeilen = _tabelle_lesen(pfad)
    if not zeilen:
        raise ImportFehler("Tabelle ist leer")
    kopf = set(zeilen[0].keys())

    def zeile(z):
        return {k: _wert(v, TABELLEN_TYPEN.get(k, "str")) for k, v in z.items() if k}

    daten = {"knoten": [], "haltungen": [], "inspektionen": [], "anlagen": []}
    if "von" in kopf and "bis" in kopf:
        typ = "haltungen"
        daten["haltungen"] = [zeile(z) for z in zeilen]
    elif "kode" in kopf and "objekt" in kopf:
        typ = "schaeden"
        gruppen = {}
        for z in map(zeile, zeilen):
            ziel_typ = "knoten" if (z.get("objektart") or "").lower().startswith(("k", "s")) else "haltung"
            schluessel = (z["objekt"], ziel_typ, z.get("datum"))
            if schluessel not in gruppen:
                gruppen[schluessel] = {"objekt": z["objekt"], "ziel_typ": ziel_typ,
                                       "datum": z.get("datum") or datetime.now().date(),
                                       "untersucher": z.get("untersucher"), "firma": z.get("firma"),
                                       "richtung": z.get("richtung") or "in Fließrichtung", "schaeden": []}
            z["kode"] = z["kode"].upper()
            if all(z.get(k) is None for k in ("klasse_d", "klasse_s", "klasse_b")):
                z["klasse_d"], z["klasse_s"], z["klasse_b"] = klassen_vorschlag(
                    z["kode"], z.get("char1"), z.get("char2"), z.get("quant1"))
            if z.get("schadensklasse") is None:
                z["schadensklasse"] = gesamtklasse(z.get("klasse_d"), z.get("klasse_s"), z.get("klasse_b"))
            gruppen[schluessel]["schaeden"].append(z)
        daten["inspektionen"] = list(gruppen.values())
    elif "anlagentyp" in kopf:
        typ = "anlagen"
        daten["anlagen"] = [zeile(z) for z in zeilen]
    elif "bezeichnung" in kopf:
        typ = "knoten"
        daten["knoten"] = [zeile(z) for z in zeilen]
    else:
        raise ImportFehler("Spalten nicht erkannt (erwartet z. B. bezeichnung; von/bis; objekt/kode; anlagentyp)")

    for eintrag in daten["knoten"] + daten["anlagen"]:
        if eintrag.get("lat") is None and eintrag.get("rechtswert") is not None:
            ll = nach_wgs84(eintrag["rechtswert"], eintrag.get("hochwert"), zone)
            if ll:
                eintrag["lat"], eintrag["lon"] = round(ll[0], 7), round(ll[1], 7)
    daten["format"] = f"Tabelle ({typ})"
    return daten


# ---------------------------------------------------------------- Einstieg

def datei_importieren(pfad, ausgeloest_von="System", anzeigename=None):
    """Importiert eine Datei und schreibt ein Protokoll. Gibt das Protokoll zurück."""
    name = anzeigename or os.path.basename(pfad)
    protokoll = ImportProtokoll(datei=name, ausgeloest_von=ausgeloest_von, erfolgreich=True, knoten=0,
                                haltungen=0, inspektionen=0, schaeden=0, anlagen=0)
    zone = current_app.config.get("UTM_ZONE", 32)
    try:
        if pfad.lower().endswith(".xml"):
            daten = isybau.lesen(pfad, zone)
        elif pfad.lower().endswith((".csv", ".xlsx")):
            daten = tabelle_lesen(pfad, zone)
        else:
            raise ImportFehler("Dateiformat nicht unterstützt (erlaubt: .xml, .csv, .xlsx)")
        protokoll.format = daten.get("format")
        meldungen = speichern(daten, name, protokoll)
        protokoll.meldungen = "\n".join(meldungen[:200]) or None
        if len(meldungen) > 200:
            protokoll.meldungen += f"\n… {len(meldungen) - 200} weitere Hinweise"
        db.session.add(protokoll)
        db.session.commit()
    except Exception as exc:  # Protokoll auch bei Fehlern schreiben
        db.session.rollback()
        protokoll = ImportProtokoll(datei=name, ausgeloest_von=ausgeloest_von, erfolgreich=False,
                                    meldungen=f"{exc}\n\n{traceback.format_exc(limit=3)}")
        db.session.add(protokoll)
        db.session.commit()
    return protokoll


def ordner_importieren(ordner=None, ausgeloest_von="Automatischer Import"):
    """Importiert alle Dateien eines Ordners (z. B. ein per OneDrive synchronisierter
    SharePoint-Ordner). Verarbeitete Dateien werden nach 'verarbeitet/', fehlerhafte
    nach 'fehler/' verschoben."""
    ordner = ordner or current_app.config["IMPORT_DIR"]
    if not os.path.isdir(ordner):
        return []
    ergebnisse = []
    for name in sorted(os.listdir(ordner)):
        pfad = os.path.join(ordner, name)
        if not os.path.isfile(pfad) or not name.lower().endswith(ENDUNGEN) or name.startswith(("~$", ".")):
            continue
        arbeit = os.path.join(ordner, ".in_arbeit_" + name)  # Endung bleibt für die Formaterkennung erhalten
        try:
            os.replace(pfad, arbeit)  # verhindert doppelte Verarbeitung
        except OSError:
            continue
        protokoll = datei_importieren(arbeit, ausgeloest_von, anzeigename=name)
        ziel_ordner = os.path.join(ordner, "verarbeitet" if protokoll.erfolgreich else "fehler")
        os.makedirs(ziel_ordner, exist_ok=True)
        stempel = datetime.now().strftime("%Y%m%d-%H%M%S")
        shutil.move(arbeit, os.path.join(ziel_ordner, f"{stempel}_{name}"))
        ergebnisse.append(protokoll)
    return ergebnisse
